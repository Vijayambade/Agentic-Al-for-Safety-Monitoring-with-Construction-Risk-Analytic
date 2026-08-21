"""
backend/routers/ai.py
---------------------
FastAPI router for General AI Construction Assistant.
Handles audio transcription, multimodal image/doc QA, and speech output.
"""
import io
import uuid
import logging
from typing import Optional
from fastapi import APIRouter, Depends, UploadFile, File, Form, HTTPException, status
from sqlalchemy.orm import Session

# Document Readers
import PyPDF2
import docx
import openpyxl

from backend.database import get_db
from backend.models.ai_history import GeneralChatHistory
from backend.routers.auth import get_current_user
from backend.models.user import User
from backend.utils.audio import transcribe_audio_file, synthesize_text_to_speech
from backend.services.ai_service import call_construction_llm

router = APIRouter(prefix="/api/v1/ai", tags=["AI General Assistant"])
logger = logging.getLogger(__name__)


def extract_text_from_file(file_name: str, file_bytes: bytes) -> str:
    """Extract text content from common document formats (PDF, DOCX, XLSX)."""
    ext = file_name.split(".")[-1].lower()
    text = ""
    
    try:
        if ext == "pdf":
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
            for page in pdf_reader.pages:
                text += page.extract_text() or ""
        elif ext in ["docx", "doc"]:
            doc = docx.Document(io.BytesIO(file_bytes))
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif ext in ["xlsx", "xls"]:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_str = " ".join([str(val) for val in row if val is not None])
                    if row_str:
                        text += row_str + "\n"
        else:
            # Fallback to general text decode
            text = file_bytes.decode("utf-8", errors="ignore")
    except Exception as e:
        logger.error("Error parsing document %s: %s", file_name, e)
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Could not parse document '{file_name}': {str(e)}"
        )
        
    return text


@router.post("/general-chat")
async def general_chat(
    session_id: str = Form(...),
    prompt: Optional[str] = Form(None),
    language: str = Form("en"),
    audio: Optional[UploadFile] = File(None),
    image: Optional[UploadFile] = File(None),
    document: Optional[UploadFile] = File(None),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Multimodal conversational endpoint for Construction AI Assistant.
    Supports audio Speech-to-Text transcription, image parsing, and document context.
    """
    user_prompt = prompt or ""
    image_bytes = None
    doc_text = None
    
    # 1. Handle Audio input (transcribe to text)
    if audio:
        try:
            audio_bytes = await audio.read()
            transcription = transcribe_audio_file(audio_bytes)
            user_prompt = transcription if not user_prompt else f"{user_prompt} {transcription}"
        except Exception as e:
            logger.error("Audio processing failed: %s", e)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=str(e)
            )

    # 2. Handle Image input
    if image:
        image_bytes = await image.read()
        
    # 3. Handle Document input
    if document:
        doc_bytes = await document.read()
        doc_text = extract_text_from_file(document.filename, doc_bytes)

    # If prompt is completely empty, throw error
    if not user_prompt.strip() and not image and not document:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Prompt text, voice audio, or file attachment is required."
        )

    # If prompt is empty but image/doc is uploaded, set default question
    if not user_prompt.strip():
        if image:
            user_prompt = "Analyze this image and identify any hazards or engineering issues."
        else:
            user_prompt = "Summarize the key engineering parameters in this document."

    # Save user message to database
    db_user_msg = GeneralChatHistory(
        session_id=session_id,
        role="user",
        message=user_prompt,
        has_audio="true" if audio else "false",
        audio_url=None
    )
    db.add(db_user_msg)
    db.commit()

    # 4. Generate response using AI Service
    ai_response = call_construction_llm(
        prompt=user_prompt,
        image_bytes=image_bytes,
        doc_text=doc_text,
        language=language
    )

    # 5. Generate Text-to-Speech voicing audio file
    tts_filename = f"{uuid.uuid4()}.mp3"
    audio_path = synthesize_text_to_speech(ai_response, tts_filename, language)

    # Save assistant response to database
    db_assistant_msg = GeneralChatHistory(
        session_id=session_id,
        role="assistant",
        message=ai_response,
        has_audio="true" if audio_path else "false",
        audio_url=audio_path if audio_path else None
    )
    db.add(db_assistant_msg)
    db.commit()

    return {
        "user_message": user_prompt,
        "response": ai_response,
        "audio_url": audio_path
    }


@router.get("/general-chat/history/{session_id}")
def get_history(
    session_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Retrieve complete dialogue log for a chat session."""
    history = (
        db.query(GeneralChatHistory)
        .filter(GeneralChatHistory.session_id == session_id)
        .order_by(GeneralChatHistory.timestamp.asc())
        .all()
    )
    return [
        {
            "role": h.role,
            "message": h.message,
            "has_audio": h.has_audio == "true",
            "audio_url": h.audio_url,
            "timestamp": h.timestamp
        }
        for h in history
    ]


@router.delete("/general-chat/history/{session_id}")
def clear_history(
    session_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Delete all dialog history for a session."""
    db.query(GeneralChatHistory).filter(GeneralChatHistory.session_id == session_id).delete()
    db.commit()
    return {"detail": "Conversation cleared successfully."}
